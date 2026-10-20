"""Portal payments report + spreadsheet export: auth, tenancy, filters, caps, secrecy."""
from __future__ import annotations

import io
from datetime import UTC, datetime

import pytest
from cryptography.fernet import Fernet
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from bims_shopify.adapters.persistence.crypto import SecretBox
from bims_shopify.adapters.persistence.models import PaymentIntentModel
from bims_shopify.adapters.persistence.tenant_repository import SqlAlchemyTenantRepository
from bims_shopify.api import portal, tenants
from bims_shopify.config import Settings
from bims_shopify.domain.tenant import Tenant

ADMIN = {"Authorization": "Bearer test-admin-token"}
SECRET_URL = "https://pay.example.com/pay?t=SIGNED-SECRET-TOKEN"


def _tenant(slug: str) -> Tenant:
    return Tenant(
        id=None, slug=slug, bims_base_url="https://bims.example.com", bims_api_key="k",
        shopify_shop_domain=f"{slug}.myshopify.com", shopify_access_token="t",
        shopify_webhook_secret="w", shopify_location_id="gid://shopify/Location/1",
        bims_posale_id=1, bims_warehouse_id=1, bims_company_id=1, bims_currency_id=1,
        bims_payment_method_id=1, default_customer_contact_id=1,
    )


def _intent(tenant_id, order_id, amount, status, day, provider="bancard_vpos"):
    return PaymentIntentModel(
        tenant_id=tenant_id, provider=provider, order_id=order_id, amount=amount,
        currency="PYG", status=status, checkout_url=SECRET_URL,
        provider_reference=f"ref-{order_id}", process_id=f"pid-{order_id}",
        created_at=datetime(2026, 10, day, 15, 0, tzinfo=UTC),
    )


@pytest.fixture
async def report_app(migrated_db):
    factory, _make_tenant = migrated_db
    settings = Settings(admin_token="test-admin-token", fernet_key=Fernet.generate_key().decode())
    app = FastAPI()
    app.state.settings = settings
    app.state.session_factory = factory
    app.include_router(tenants.router)
    app.include_router(portal.router)
    async with factory() as session:
        repo = SqlAlchemyTenantRepository(session, SecretBox(settings.fernet_key))
        acme = await repo.create(_tenant("acme"))
        other = await repo.create(_tenant("other"))
        session.add_all(
            [
                _intent(acme.id, "A1", 1234567.6, "confirmed", 1),
                _intent(acme.id, "A2", 50000, "pending", 1),
                _intent(acme.id, "=HYPERLINK(1)", 25000, "confirmed", 2),
                _intent(acme.id, "A4", 10000, "failed", 3, provider="pagopar"),
                _intent(other.id, "O1", 999999, "confirmed", 1),
            ]
        )
        await session.commit()
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        yield client, app


RANGE = {"date_from": "2026-10-01", "date_to": "2026-10-31"}


async def test_report_requires_auth(report_app):
    client, _ = report_app
    assert (await client.get("/api/portal/acme/payments/report")).status_code == 401
    assert (await client.get("/api/portal/acme/payments/export")).status_code == 401


async def test_report_returns_rows_and_aggregates_without_secrets(report_app):
    client, _ = report_app
    response = await client.get("/api/portal/acme/payments/report", params=RANGE, headers=ADMIN)
    assert response.status_code == 200
    body = response.json()
    assert body["total"] == 4
    assert {r["order_id"] for r in body["rows"]} == {"A1", "A2", "=HYPERLINK(1)", "A4"}
    by_status = {s["status"]: s for s in body["by_status"]}
    assert by_status["confirmed"]["count"] == 2
    assert [d["date"] for d in body["by_day"]] == ["2026-10-01", "2026-10-02", "2026-10-03"]
    assert body["filters"]["date_from"] == "2026-10-01"
    assert "SIGNED-SECRET-TOKEN" not in response.text
    assert "checkout_url" not in response.text


async def test_report_never_leaks_another_tenants_payments(report_app):
    client, _ = report_app
    acme = (await client.get("/api/portal/acme/payments/report", params=RANGE, headers=ADMIN)).json()
    other = (await client.get("/api/portal/other/payments/report", params=RANGE, headers=ADMIN)).json()
    assert "O1" not in {r["order_id"] for r in acme["rows"]}
    assert {r["order_id"] for r in other["rows"]} == {"O1"}
    assert other["by_status"] == [{"status": "confirmed", "count": 1, "total_amount": 999999}]


async def test_other_tenants_portal_token_is_rejected(report_app):
    client, _ = report_app
    minted = await client.post("/tenants/other/portal-token", headers=ADMIN)
    token = minted.json()["portal_token"]
    ok = await client.get(
        "/api/portal/other/payments/report", headers={"Authorization": f"Bearer {token}"}
    )
    denied = await client.get(
        "/api/portal/acme/payments/report", headers={"Authorization": f"Bearer {token}"}
    )
    assert ok.status_code == 200
    assert denied.status_code == 401


async def test_report_filters_and_pagination(report_app):
    client, _ = report_app
    confirmed = (
        await client.get(
            "/api/portal/acme/payments/report", params={**RANGE, "status": "confirmed"}, headers=ADMIN
        )
    ).json()
    assert confirmed["total"] == 2
    page = (
        await client.get(
            "/api/portal/acme/payments/report", params={**RANGE, "limit": 1, "offset": 1}, headers=ADMIN
        )
    ).json()
    assert len(page["rows"]) == 1 and page["total"] == 4 and page["offset"] == 1


@pytest.mark.parametrize(
    "params",
    [
        {"date_from": "2026-10-31", "date_to": "2026-10-01"},
        {"date_from": "2024-01-01", "date_to": "2026-10-31"},
        {"date_from": "not-a-date"},
        {"limit": 5000},
    ],
)
async def test_report_rejects_invalid_parameters(report_app, params):
    client, _ = report_app
    response = await client.get("/api/portal/acme/payments/report", params=params, headers=ADMIN)
    assert response.status_code == 422


async def test_export_is_an_xlsx_with_safe_cells_and_no_secrets(report_app):
    client, _ = report_app
    response = await client.get("/api/portal/acme/payments/export", params=RANGE, headers=ADMIN)
    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"] == (
        'attachment; filename="payments-acme-2026-10-01_2026-10-31.xlsx"'
    )
    assert b"SIGNED-SECRET-TOKEN" not in response.content
    sheet = load_workbook(io.BytesIO(response.content))["Payments"]
    header = [c.value for c in sheet[1]]
    assert "checkout_url" not in header
    assert header[:3] == ["id", "order_id", "created_at"]
    rows = {r[1].value: r for r in sheet.iter_rows(min_row=2)}
    amount_col = header.index("amount")
    big = rows["A1"][amount_col]
    assert big.value == 1234568 and big.number_format == "#,##0"
    assert rows["A1"][2].value == "2026-10-01T15:00:00+00:00"
    injected = rows["'=HYPERLINK(1)"][1]
    assert injected.data_type == "s"
    assert "O1" not in rows


async def test_export_is_row_capped_and_flags_truncation(report_app, monkeypatch):
    client, _ = report_app
    monkeypatch.setattr(portal, "EXPORT_ROW_CAP", 2)
    response = await client.get("/api/portal/acme/payments/export", params=RANGE, headers=ADMIN)
    assert response.headers["x-report-truncated"] == "true"
    sheet = load_workbook(io.BytesIO(response.content))["Payments"]
    assert sheet.max_row == 3

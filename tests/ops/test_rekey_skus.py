"""Decision-table tests for the rekey_skus ops script.

Covers every branch of the SKU-classification logic (empty_sku,
already_keyed, planned_rewrite, name_mismatch, unresolved), idempotency of a
second scan after a rewrite, and product-grouping of bulk-update calls.
"""
from __future__ import annotations

import json
import os
import tempfile
from datetime import UTC, datetime
from pathlib import Path

import httpx
import pytest
import respx
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from bims_shopify.adapters.bims.client import BIMSClient
from bims_shopify.adapters.persistence.models import RekeyReportModel
from bims_shopify.adapters.shopify.client import ShopifyClient
from bims_shopify.ops.rekey_skus import (
    ScanResult,
    apply_rewrites,
    build_report_workbook,
    partition_duplicate_targets,
    persist_rekey_report,
    scan,
    write_report_xlsx,
)

GRAPHQL_URL = "https://acme.myshopify.com/admin/api/2025-07/graphql.json/"
BIMS_URL = "https://bims.example.com"


def _variants_page(nodes: list[dict], has_next: bool = False, end_cursor: str | None = None) -> dict:
    return {
        "data": {
            "productVariants": {
                "pageInfo": {"hasNextPage": has_next, "endCursor": end_cursor},
                "nodes": nodes,
            }
        }
    }


def _bulk_update_result(variant_ids: list[str]) -> dict:
    return {
        "data": {
            "productVariantsBulkUpdate": {
                "productVariants": [{"id": vid} for vid in variant_ids],
                "userErrors": [],
            }
        }
    }


def _mock_graphql_variants(nodes: list[dict]) -> None:
    respx.post(GRAPHQL_URL).mock(return_value=httpx.Response(200, json=_variants_page(nodes)))


def _mock_company6_index(code2_values: list[str]) -> None:
    items = [{"code2": value} for value in code2_values]
    respx.get(f"{BIMS_URL}/api/products/index.json").mock(
        return_value=httpx.Response(200, json={"status": "ok", "data": items})
    )


def _mock_view_lookup(products_by_sku: dict[str, dict]) -> None:
    def responder(request: httpx.Request) -> httpx.Response:
        sku = request.url.params.get("id")
        product = products_by_sku.get(sku)
        if product is None:
            return httpx.Response(200, json={"status": "error", "code": "404", "message": "not found"})
        return httpx.Response(200, json={"status": "ok", "data": product})

    respx.get(f"{BIMS_URL}/api/products/view.json").mock(side_effect=responder)


def _variant(variant_id: str, sku: str, product_id: str, product_title: str) -> dict:
    return {"id": variant_id, "sku": sku, "product": {"id": product_id, "title": product_title}}


@respx.mock
async def test_empty_sku_is_categorized_as_empty_sku(tenant):
    _mock_graphql_variants(
        [_variant("gid://shopify/ProductVariant/1", "  ", "gid://shopify/Product/1", "Widget")]
    )
    _mock_company6_index([])
    _mock_view_lookup({})

    shopify_client = ShopifyClient(tenant)
    bims_client = BIMSClient(tenant)
    result = await scan(shopify_client, bims_client)

    assert result.total_variants == 1
    assert result.empty_sku == 1
    assert result.already_keyed == 0
    assert result.planned_rewrites == []
    assert result.name_mismatch == []
    assert result.unresolved == []


@respx.mock
async def test_sku_already_in_company6_set_is_already_keyed(tenant):
    _mock_graphql_variants(
        [_variant("gid://shopify/ProductVariant/1", "SKU-100", "gid://shopify/Product/1", "Widget")]
    )
    _mock_company6_index(["SKU-100"])
    _mock_view_lookup({})

    shopify_client = ShopifyClient(tenant)
    bims_client = BIMSClient(tenant)
    result = await scan(shopify_client, bims_client)

    assert result.empty_sku == 0
    assert result.already_keyed == 1
    assert result.planned_rewrites == []


@respx.mock
async def test_resolved_via_bims_view_with_matching_names_is_planned_rewrite(tenant):
    _mock_graphql_variants(
        [_variant("gid://shopify/ProductVariant/1", "12345", "gid://shopify/Product/1", "Blue Widget (L)")]
    )
    _mock_company6_index(["SKU-999"])
    _mock_view_lookup({"12345": {"code2": "SKU-999", "name": "Blue Widget"}})

    shopify_client = ShopifyClient(tenant)
    bims_client = BIMSClient(tenant)
    result = await scan(shopify_client, bims_client)

    assert result.already_keyed == 0
    assert result.name_mismatch == []
    assert len(result.planned_rewrites) == 1
    rewrite = result.planned_rewrites[0]
    assert rewrite["old_sku"] == "12345"
    assert rewrite["new_sku"] == "SKU-999"
    assert rewrite["variant_id"] == "gid://shopify/ProductVariant/1"
    assert rewrite["product_id"] == "gid://shopify/Product/1"


@respx.mock
async def test_resolved_via_bims_view_with_mismatching_names_is_name_mismatch(tenant):
    _mock_graphql_variants(
        [_variant("gid://shopify/ProductVariant/1", "12345", "gid://shopify/Product/1", "Red Sneakers")]
    )
    _mock_company6_index(["SKU-999"])
    _mock_view_lookup({"12345": {"code2": "SKU-999", "name": "Blue Widget"}})

    shopify_client = ShopifyClient(tenant)
    bims_client = BIMSClient(tenant)
    result = await scan(shopify_client, bims_client)

    assert result.planned_rewrites == []
    assert len(result.name_mismatch) == 1
    mismatch = result.name_mismatch[0]
    assert mismatch["old_sku"] == "12345"
    assert mismatch["new_sku"] == "SKU-999"
    assert mismatch["bims_name"] == "Blue Widget"


@respx.mock
async def test_view_lookup_not_found_is_unresolved(tenant):
    _mock_graphql_variants(
        [_variant("gid://shopify/ProductVariant/1", "unknown-id", "gid://shopify/Product/1", "Widget")]
    )
    _mock_company6_index(["SKU-999"])
    _mock_view_lookup({})

    shopify_client = ShopifyClient(tenant)
    bims_client = BIMSClient(tenant)
    result = await scan(shopify_client, bims_client)

    assert result.planned_rewrites == []
    assert result.name_mismatch == []
    assert len(result.unresolved) == 1
    assert result.unresolved[0]["sku"] == "unknown-id"


@respx.mock
async def test_view_lookup_code2_not_in_company6_set_is_unresolved(tenant):
    _mock_graphql_variants(
        [_variant("gid://shopify/ProductVariant/1", "12345", "gid://shopify/Product/1", "Widget")]
    )
    _mock_company6_index(["SKU-999"])
    _mock_view_lookup({"12345": {"code2": "SKU-000", "name": "Widget"}})

    shopify_client = ShopifyClient(tenant)
    bims_client = BIMSClient(tenant)
    result = await scan(shopify_client, bims_client)

    assert result.planned_rewrites == []
    assert result.name_mismatch == []
    assert len(result.unresolved) == 1


@respx.mock
async def test_second_scan_after_apply_reports_zero_planned_rewrites_for_rewritten_sku(tenant):
    """Idempotency: once a variant's SKU equals its company-6 code2, it is already_keyed."""
    _mock_graphql_variants(
        [_variant("gid://shopify/ProductVariant/1", "SKU-999", "gid://shopify/Product/1", "Blue Widget")]
    )
    _mock_company6_index(["SKU-999"])
    _mock_view_lookup({})

    shopify_client = ShopifyClient(tenant)
    bims_client = BIMSClient(tenant)
    result = await scan(shopify_client, bims_client)

    assert result.already_keyed == 1
    assert result.planned_rewrites == []


@respx.mock
async def test_apply_rewrites_groups_variants_by_product_into_one_call_each(tenant):
    calls: list[dict] = []

    def responder(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        if "ProductVariantsBulkUpdate" in body["query"]:
            calls.append(body["variables"])
            variant_ids = [v["id"] for v in body["variables"]["variants"]]
            return httpx.Response(200, json=_bulk_update_result(variant_ids))
        raise AssertionError(f"unexpected GraphQL operation: {body['query']}")

    respx.post(GRAPHQL_URL).mock(side_effect=responder)

    shopify_client = ShopifyClient(tenant)
    planned_rewrites = [
        {
            "variant_id": "gid://shopify/ProductVariant/1",
            "product_id": "gid://shopify/Product/1",
            "old_sku": "111",
            "new_sku": "SKU-A",
        },
        {
            "variant_id": "gid://shopify/ProductVariant/2",
            "product_id": "gid://shopify/Product/1",
            "old_sku": "222",
            "new_sku": "SKU-B",
        },
        {
            "variant_id": "gid://shopify/ProductVariant/3",
            "product_id": "gid://shopify/Product/2",
            "old_sku": "333",
            "new_sku": "SKU-C",
        },
    ]

    await apply_rewrites(shopify_client, planned_rewrites)

    assert len(calls) == 2
    by_product = {call["productId"]: call["variants"] for call in calls}
    assert {v["id"] for v in by_product["gid://shopify/Product/1"]} == {
        "gid://shopify/ProductVariant/1",
        "gid://shopify/ProductVariant/2",
    }
    assert by_product["gid://shopify/Product/2"] == [
        {"id": "gid://shopify/ProductVariant/3", "sku": "SKU-C"}
    ]


@respx.mock
async def test_transient_lookup_failure_is_unresolved_and_does_not_abort_scan(tenant):
    """A single lookup raising a transient network error must not kill the whole scan.

    Regression test: previously only BIMSAPIError/HTTPStatusError/TimeoutError were
    caught, so httpx.ConnectError (a realistic failure at concurrency 5 against a
    live network) would propagate through asyncio.gather and abort the entire scan.
    """
    _mock_graphql_variants(
        [
            _variant("gid://shopify/ProductVariant/1", "broken-lookup", "gid://shopify/Product/1", "Widget"),
            _variant("gid://shopify/ProductVariant/2", "12345", "gid://shopify/Product/2", "Blue Widget"),
        ]
    )
    _mock_company6_index(["SKU-999"])

    def responder(request: httpx.Request) -> httpx.Response:
        sku = request.url.params.get("id")
        if sku == "broken-lookup":
            raise httpx.ConnectError("connection refused", request=request)
        if sku == "12345":
            return httpx.Response(200, json={"status": "ok", "data": {"Product": {"code2": "SKU-999", "name": "Blue Widget"}}})
        return httpx.Response(200, json={"status": "error", "code": "404", "message": "not found"})

    respx.get(f"{BIMS_URL}/api/products/view.json").mock(side_effect=responder)

    shopify_client = ShopifyClient(tenant)
    bims_client = BIMSClient(tenant)
    result = await scan(shopify_client, bims_client)

    assert len(result.unresolved) == 1
    assert result.unresolved[0]["sku"] == "broken-lookup"
    assert "ConnectError" in result.unresolved[0]["error"]
    assert len(result.planned_rewrites) == 1
    assert result.planned_rewrites[0]["old_sku"] == "12345"


@respx.mock
async def test_apply_rewrites_isolates_per_product_failures_and_continues(tenant):
    """One product's bulk update failing (userErrors) must not abort remaining products."""

    def responder(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        variables = body["variables"]
        if variables["productId"] == "gid://shopify/Product/1":
            return httpx.Response(
                200,
                json={
                    "data": {
                        "productVariantsBulkUpdate": {
                            "productVariants": [],
                            "userErrors": [{"field": ["sku"], "message": "SKU already taken"}],
                        }
                    }
                },
            )
        variant_ids = [v["id"] for v in variables["variants"]]
        return httpx.Response(200, json=_bulk_update_result(variant_ids))

    respx.post(GRAPHQL_URL).mock(side_effect=responder)

    shopify_client = ShopifyClient(tenant)
    planned_rewrites = [
        {
            "variant_id": "gid://shopify/ProductVariant/1",
            "product_id": "gid://shopify/Product/1",
            "old_sku": "111",
            "new_sku": "SKU-A",
        },
        {
            "variant_id": "gid://shopify/ProductVariant/3",
            "product_id": "gid://shopify/Product/2",
            "old_sku": "333",
            "new_sku": "SKU-C",
        },
    ]

    result = await apply_rewrites(shopify_client, planned_rewrites)

    assert len(result.failed) == 1
    assert result.failed[0]["product_id"] == "gid://shopify/Product/1"
    assert "SKU already taken" in result.failed[0]["error"]
    assert len(result.succeeded) == 1
    assert result.succeeded[0]["product_id"] == "gid://shopify/Product/2"


def test_partition_duplicate_targets_moves_all_conflicting_variants_out():
    planned_rewrites = [
        {"variant_id": "v1", "product_id": "p1", "old_sku": "a", "new_sku": "SKU-DUP"},
        {"variant_id": "v2", "product_id": "p2", "old_sku": "b", "new_sku": "SKU-DUP"},
        {"variant_id": "v3", "product_id": "p3", "old_sku": "c", "new_sku": "SKU-UNIQUE"},
    ]

    safe, duplicates = partition_duplicate_targets(planned_rewrites)

    assert safe == [{"variant_id": "v3", "product_id": "p3", "old_sku": "c", "new_sku": "SKU-UNIQUE"}]
    assert {d["variant_id"] for d in duplicates} == {"v1", "v2"}


@respx.mock
async def test_scan_moves_duplicate_new_sku_targets_out_of_planned_rewrites(tenant):
    """Two variants resolving to the same code2 must not both be queued for rewrite."""
    _mock_graphql_variants(
        [
            _variant("gid://shopify/ProductVariant/1", "111", "gid://shopify/Product/1", "Blue Widget"),
            _variant("gid://shopify/ProductVariant/2", "222", "gid://shopify/Product/2", "Blue Widget"),
        ]
    )
    _mock_company6_index(["SKU-999"])
    _mock_view_lookup(
        {
            "111": {"code2": "SKU-999", "name": "Blue Widget"},
            "222": {"code2": "SKU-999", "name": "Blue Widget"},
        }
    )

    shopify_client = ShopifyClient(tenant)
    bims_client = BIMSClient(tenant)
    result = await scan(shopify_client, bims_client)

    assert result.planned_rewrites == []
    assert len(result.duplicate_target) == 2
    assert {d["old_sku"] for d in result.duplicate_target} == {"111", "222"}


@respx.mock
async def test_shopify_variant_pagination_yields_every_variant_across_pages(tenant):
    """iter_all_variants must not drop or duplicate entries when Shopify paginates."""
    page1 = _variants_page(
        [_variant("gid://shopify/ProductVariant/1", "SKU-1", "gid://shopify/Product/1", "Widget One")],
        has_next=True,
        end_cursor="cursor-1",
    )
    page2 = _variants_page(
        [_variant("gid://shopify/ProductVariant/2", "SKU-2", "gid://shopify/Product/2", "Widget Two")],
        has_next=False,
    )

    calls = {"count": 0}

    def responder(request: httpx.Request) -> httpx.Response:
        calls["count"] += 1
        return httpx.Response(200, json=page1 if calls["count"] == 1 else page2)

    respx.post(GRAPHQL_URL).mock(side_effect=responder)
    _mock_company6_index(["SKU-1", "SKU-2"])
    _mock_view_lookup({})

    shopify_client = ShopifyClient(tenant)
    bims_client = BIMSClient(tenant)
    result = await scan(shopify_client, bims_client)

    assert result.total_variants == 2
    assert result.already_keyed == 2
    assert calls["count"] == 2


def _fixture_scan_result(*, duplicates: int, unresolved: int, name_mismatch: int) -> ScanResult:
    result = ScanResult(total_variants=100, empty_sku=5, already_keyed=10)
    for i in range(duplicates):
        product_title = "Blue Widget" if i % 2 == 0 else "Ant Widget"
        new_sku = "SKU-DUP-A" if i % 2 == 0 else "SKU-DUP-B"
        result.duplicate_target.append(
            {
                "variant_id": f"gid://shopify/ProductVariant/dup{i}",
                "product_id": f"gid://shopify/Product/dup{i}",
                "product_title": product_title,
                "old_sku": f"old-dup-{i}",
                "new_sku": new_sku,
                "bims_name": f"Bims Name {i}",
            }
        )
    for i in range(unresolved):
        result.unresolved.append(
            {
                "variant_id": f"gid://shopify/ProductVariant/unres{i}",
                "product_title": f"Unresolved Product {i}",
                "sku": f"unresolved-sku-{i}",
            }
        )
    for i in range(name_mismatch):
        result.name_mismatch.append(
            {
                "variant_id": f"gid://shopify/ProductVariant/mismatch{i}",
                "product_id": f"gid://shopify/Product/mismatch{i}",
                "product_title": f"Shopify Title {i}",
                "old_sku": f"old-mismatch-{i}",
                "new_sku": f"SKU-MISMATCH-{i}",
                "bims_name": f"Bims Name {i}",
            }
        )
    return result


def test_build_report_workbook_has_exactly_four_sheets():
    result = _fixture_scan_result(duplicates=15, unresolved=12, name_mismatch=12)
    workbook = build_report_workbook(
        result,
        shop_domain="acme.myshopify.com",
        tenant_slug="acme",
        scanned_at=datetime(2026, 1, 1, tzinfo=UTC),
    )

    assert workbook.sheetnames == ["Summary", "Duplicates", "Unresolved", "Name mismatch"]


def test_build_report_workbook_does_not_truncate_full_lists():
    result = _fixture_scan_result(duplicates=15, unresolved=12, name_mismatch=12)
    workbook = build_report_workbook(
        result,
        shop_domain="acme.myshopify.com",
        tenant_slug="acme",
        scanned_at=datetime(2026, 1, 1, tzinfo=UTC),
    )

    # header row + one row per item, no 10-item preview truncation.
    assert workbook["Duplicates"].max_row - 1 == 15
    assert workbook["Unresolved"].max_row - 1 == 12
    assert workbook["Name mismatch"].max_row - 1 == 12


def test_build_report_workbook_duplicates_sheet_header_bold_and_frozen():
    result = _fixture_scan_result(duplicates=2, unresolved=0, name_mismatch=0)
    workbook = build_report_workbook(
        result,
        shop_domain="acme.myshopify.com",
        tenant_slug="acme",
        scanned_at=datetime(2026, 1, 1, tzinfo=UTC),
    )
    ws = workbook["Duplicates"]

    assert [cell.value for cell in ws[1]] == [
        "product_title",
        "size/bims_name",
        "old_sku",
        "target_code2(new_sku)",
        "variant_id",
        "product_id",
        "KEEP? (YES/NO)",
    ]
    assert all(cell.font.bold for cell in ws[1])
    assert ws.freeze_panes == "A2"


def test_build_report_workbook_name_mismatch_sheet_header_order():
    result = _fixture_scan_result(duplicates=0, unresolved=0, name_mismatch=1)
    workbook = build_report_workbook(
        result,
        shop_domain="acme.myshopify.com",
        tenant_slug="acme",
        scanned_at=datetime(2026, 1, 1, tzinfo=UTC),
    )
    ws = workbook["Name mismatch"]

    assert [cell.value for cell in ws[1]] == [
        "product_title",
        "bims_name",
        "old_sku",
        "proposed_code2",
        "variant_id",
        "APPROVE? (YES/NO)",
    ]
    assert ws.freeze_panes == "A2"
    assert all(cell.font.bold for cell in ws[1])


def test_build_report_workbook_duplicates_sheet_grouped_by_product_title_and_new_sku():
    result = _fixture_scan_result(duplicates=4, unresolved=0, name_mismatch=0)
    workbook = build_report_workbook(
        result,
        shop_domain="acme.myshopify.com",
        tenant_slug="acme",
        scanned_at=datetime(2026, 1, 1, tzinfo=UTC),
    )
    ws = workbook["Duplicates"]

    rows = [(row[0].value, row[3].value) for row in ws.iter_rows(min_row=2)]
    assert rows == sorted(rows)


def test_write_report_xlsx_round_trips_via_tmp_path(tmp_path):
    result = _fixture_scan_result(duplicates=1, unresolved=1, name_mismatch=1)
    path = tmp_path / "report.xlsx"

    write_report_xlsx(
        result,
        str(path),
        shop_domain="acme.myshopify.com",
        tenant_slug="acme",
        scanned_at=datetime(2026, 1, 1, tzinfo=UTC),
    )

    assert path.exists()
    workbook = load_workbook(str(path))
    assert workbook.sheetnames == ["Summary", "Duplicates", "Unresolved", "Name mismatch"]


@pytest.fixture
async def rekey_reports_session_factory(monkeypatch):
    with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as db_file:
        pass
    database_url = f"sqlite+aiosqlite:///{db_file.name}"

    from alembic import command
    from alembic.config import Config

    repo_root = Path(__file__).resolve().parents[2]
    cfg = Config(str(repo_root / "alembic.ini"))
    cfg.set_main_option("script_location", str(repo_root / "alembic"))
    os.environ["ALEMBIC_DATABASE_URL"] = database_url
    try:
        await __import__("asyncio").to_thread(command.upgrade, cfg, "head")
    finally:
        os.environ.pop("ALEMBIC_DATABASE_URL", None)

    from bims_shopify.config import Settings

    monkeypatch.setattr(
        "bims_shopify.ops.rekey_skus.get_settings",
        lambda: Settings(database_url=database_url),
    )

    engine = create_async_engine(database_url, future=True)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        yield factory
    finally:
        await engine.dispose()
        Path(db_file.name).unlink(missing_ok=True)


async def test_persist_rekey_report_writes_row_with_tenant_id_and_payload(
    rekey_reports_session_factory,
):
    payload = {"total_variants": 3, "duplicate_target": []}

    await persist_rekey_report(7, payload)

    async with rekey_reports_session_factory() as session:
        rows = (await session.execute(select(RekeyReportModel))).scalars().all()

    assert len(rows) == 1
    assert rows[0].tenant_id == 7
    assert rows[0].payload == payload


@respx.mock
async def test_run_with_report_xlsx_includes_report_path_in_json_summary(tmp_path, monkeypatch):
    monkeypatch.setenv("SHOPIFY_TOKEN", "shpat_test")
    monkeypatch.setenv("BIMS_KEY", "bims_secret")

    _mock_graphql_variants(
        [_variant("gid://shopify/ProductVariant/1", "SKU-100", "gid://shopify/Product/1", "Widget")]
    )
    _mock_company6_index(["SKU-100"])
    _mock_view_lookup({})

    from bims_shopify.ops import rekey_skus as rekey_skus_module

    report_path = tmp_path / "report.xlsx"
    parser = rekey_skus_module._build_arg_parser()
    args = parser.parse_args(
        [
            "--standalone",
            "--shop",
            "acme.myshopify.com",
            "--shopify-token-env",
            "SHOPIFY_TOKEN",
            "--bims-key-env",
            "BIMS_KEY",
            "--bims-url",
            BIMS_URL,
            "--report-xlsx",
            str(report_path),
        ]
    )

    summary = await rekey_skus_module._run(args)

    assert summary["report_path"] == str(report_path)
    assert report_path.exists()


@respx.mock
async def test_run_standalone_does_not_persist_rekey_report(tmp_path, monkeypatch):
    monkeypatch.setenv("SHOPIFY_TOKEN", "shpat_test")
    monkeypatch.setenv("BIMS_KEY", "bims_secret")

    _mock_graphql_variants(
        [_variant("gid://shopify/ProductVariant/1", "SKU-100", "gid://shopify/Product/1", "Widget")]
    )
    _mock_company6_index(["SKU-100"])
    _mock_view_lookup({})

    from bims_shopify.ops import rekey_skus as rekey_skus_module

    async def _fail_if_called(*args, **kwargs):
        raise AssertionError("persist_rekey_report must not be called in standalone mode")

    monkeypatch.setattr(rekey_skus_module, "persist_rekey_report", _fail_if_called)

    parser = rekey_skus_module._build_arg_parser()
    args = parser.parse_args(
        [
            "--standalone",
            "--shop",
            "acme.myshopify.com",
            "--shopify-token-env",
            "SHOPIFY_TOKEN",
            "--bims-key-env",
            "BIMS_KEY",
            "--bims-url",
            BIMS_URL,
        ]
    )

    await rekey_skus_module._run(args)


@respx.mock
async def test_run_db_mode_persists_rekey_report(tenant, monkeypatch):
    _mock_graphql_variants(
        [_variant("gid://shopify/ProductVariant/1", "SKU-100", "gid://shopify/Product/1", "Widget")]
    )
    _mock_company6_index(["SKU-100"])
    _mock_view_lookup({})

    from bims_shopify.ops import rekey_skus as rekey_skus_module

    async def _fake_load_tenant_from_db(tenant_slug: str):
        return tenant

    calls: list[tuple[int, dict]] = []

    async def _fake_persist(tenant_id: int, payload: dict) -> None:
        calls.append((tenant_id, payload))

    monkeypatch.setattr(rekey_skus_module, "_load_tenant_from_db", _fake_load_tenant_from_db)
    monkeypatch.setattr(rekey_skus_module, "persist_rekey_report", _fake_persist)

    parser = rekey_skus_module._build_arg_parser()
    args = parser.parse_args(["acme"])

    await rekey_skus_module._run(args)

    assert len(calls) == 1
    assert calls[0][0] == tenant.id


@respx.mock
async def test_bims_company6_index_pagination_collects_code2_across_pages(tenant):
    """fetch_company_code2_set must page through the full BIMS index, not just page 1."""
    _mock_graphql_variants(
        [_variant("gid://shopify/ProductVariant/1", "SKU-PAGE2", "gid://shopify/Product/1", "Widget")]
    )

    from bims_shopify.ops import rekey_skus as rekey_skus_module

    page_limit = 2
    monkey_limit = rekey_skus_module.BIMS_INDEX_PAGE_LIMIT
    rekey_skus_module.BIMS_INDEX_PAGE_LIMIT = page_limit
    try:
        page_one_items = [{"code2": f"SKU-{i}"} for i in range(page_limit)]
        page_two_items = [{"code2": "SKU-PAGE2"}]

        def responder(request: httpx.Request) -> httpx.Response:
            offset = int(request.url.params.get("offset"))
            items = page_one_items if offset == 0 else page_two_items
            return httpx.Response(200, json={"status": "ok", "data": items})

        respx.get(f"{BIMS_URL}/api/products/index.json").mock(side_effect=responder)
        _mock_view_lookup({})

        shopify_client = ShopifyClient(tenant)
        bims_client = BIMSClient(tenant)
        result = await scan(shopify_client, bims_client)
    finally:
        rekey_skus_module.BIMS_INDEX_PAGE_LIMIT = monkey_limit

    assert result.already_keyed == 1
    assert result.unresolved == []

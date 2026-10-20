"""Render the payments report as an XLSX workbook.

Amounts are written as numbers (so the spreadsheet can sum them) with a
no-decimals format, matching how guaraníes are displayed. Dates stay ISO
8601 text. Text cells that a spreadsheet would interpret as a formula are
neutralised. ``checkout_url`` is never part of the row data it receives.
"""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

PYG_FORMAT = "#,##0"
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

_COLUMNS = (
    "id",
    "order_id",
    "created_at",
    "updated_at",
    "provider",
    "status",
    "amount",
    "currency",
    "provider_reference",
    "process_id",
)


def _safe_text(value: Any) -> Any:
    """Prefix text that would be parsed as a formula so it stays inert."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _amount(value: float | None) -> int:
    return round(value or 0)


def build_payments_workbook(
    rows: list[dict[str, Any]],
    by_status: list[dict[str, Any]],
    by_day: list[dict[str, Any]],
    *,
    truncated: bool,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Payments"
    sheet.append(list(_COLUMNS))
    for row in rows:
        sheet.append([_safe_text(row.get(name)) for name in _COLUMNS])
    amount_col = _COLUMNS.index("amount") + 1
    for index in range(2, len(rows) + 2):
        cell = sheet.cell(row=index, column=amount_col)
        cell.value = _amount(rows[index - 2].get("amount"))
        cell.number_format = PYG_FORMAT
    _bold_header(sheet)

    summary = workbook.create_sheet("Summary")
    summary.append(["status", "count", "total_amount"])
    for item in by_status:
        summary.append([_safe_text(item["status"]), item["count"], _amount(item["total_amount"])])
    summary.append([])
    summary.append(["date", "count", "total_amount"])
    for item in by_day:
        summary.append([item["date"], item["count"], _amount(item["total_amount"])])
    for line in summary.iter_rows():
        line[-1].number_format = PYG_FORMAT
    summary["A1"].font = Font(bold=True)
    if truncated:
        summary.append([])
        summary.append(["note", "The Payments sheet was truncated to the export row limit."])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _bold_header(sheet) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
